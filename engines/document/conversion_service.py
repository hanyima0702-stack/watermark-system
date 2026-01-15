"""
Document format conversion service using LibreOffice headless mode.
Supports batch conversion and quality checking.
"""
import os
import subprocess
import tempfile
import shutil
from pathlib import Path
from typing import List, Dict, Any, Optional, Tuple
import logging
import time
import threading
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from enum import Enum

from .base_processor import DocumentProcessingError

logger = logging.getLogger(__name__)


class ConversionFormat(Enum):
    """Supported conversion formats."""
    PDF = "pdf"
    DOCX = "docx"
    XLSX = "xlsx"
    PPTX = "pptx"
    ODT = "odt"
    ODS = "ods"
    ODP = "odp"
    RTF = "rtf"
    TXT = "txt"
    HTML = "html"


@dataclass
class ConversionJob:
    """Represents a document conversion job."""
    input_path: Path
    output_path: Path
    target_format: ConversionFormat
    options: Dict[str, Any]
    status: str = "pending"
    error_message: Optional[str] = None
    start_time: Optional[float] = None
    end_time: Optional[float] = None
    
    @property
    def duration(self) -> Optional[float]:
        """Get conversion duration in seconds."""
        if self.start_time and self.end_time:
            return self.end_time - self.start_time
        return None


class ConversionQualityChecker:
    """Checks conversion quality and validates output."""
    
    def __init__(self):
        self.quality_thresholds = {
            'min_file_size_ratio': 0.1,  # Output should be at least 10% of input size
            'max_file_size_ratio': 10.0,  # Output shouldn't be more than 10x input size
        }
    
    def check_conversion_quality(self, input_path: Path, output_path: Path, 
                               target_format: ConversionFormat) -> Dict[str, Any]:
        """Check the quality of a conversion."""
        quality_report = {
            'is_valid': True,
            'issues': [],
            'metrics': {}
        }
        
        try:
            # Check if output file exists
            if not output_path.exists():
                quality_report['is_valid'] = False
                quality_report['issues'].append("Output file does not exist")
                return quality_report
            
            # Check file sizes
            input_size = input_path.stat().st_size
            output_size = output_path.stat().st_size
            
            quality_report['metrics']['input_size'] = input_size
            quality_report['metrics']['output_size'] = output_size
            quality_report['metrics']['size_ratio'] = output_size / input_size if input_size > 0 else 0
            
            # Validate file size ratio
            size_ratio = quality_report['metrics']['size_ratio']
            if size_ratio < self.quality_thresholds['min_file_size_ratio']:
                quality_report['issues'].append(f"Output file too small (ratio: {size_ratio:.2f})")
            elif size_ratio > self.quality_thresholds['max_file_size_ratio']:
                quality_report['issues'].append(f"Output file too large (ratio: {size_ratio:.2f})")
            
            # Check if file is readable
            try:
                if target_format == ConversionFormat.PDF:
                    self._validate_pdf(output_path, quality_report)
                elif target_format in [ConversionFormat.DOCX, ConversionFormat.XLSX, ConversionFormat.PPTX]:
                    self._validate_office_document(output_path, quality_report)
                elif target_format == ConversionFormat.TXT:
                    self._validate_text_file(output_path, quality_report)
                    
            except Exception as e:
                quality_report['issues'].append(f"File validation failed: {str(e)}")
            
            # Set overall validity
            quality_report['is_valid'] = len(quality_report['issues']) == 0
            
        except Exception as e:
            quality_report['is_valid'] = False
            quality_report['issues'].append(f"Quality check failed: {str(e)}")
            
        return quality_report
    
    def _validate_pdf(self, pdf_path: Path, quality_report: Dict[str, Any]):
        """Validate PDF file."""
        try:
            import fitz  # PyMuPDF
            doc = fitz.open(str(pdf_path))
            page_count = len(doc)
            quality_report['metrics']['page_count'] = page_count
            
            if page_count == 0:
                quality_report['issues'].append("PDF has no pages")
            
            doc.close()
            
        except Exception as e:
            quality_report['issues'].append(f"PDF validation failed: {str(e)}")
    
    def _validate_office_document(self, doc_path: Path, quality_report: Dict[str, Any]):
        """Validate Office document."""
        try:
            # Try to open with appropriate library
            if doc_path.suffix.lower() == '.docx':
                from docx import Document
                doc = Document(str(doc_path))
                quality_report['metrics']['paragraph_count'] = len(doc.paragraphs)
                
            elif doc_path.suffix.lower() == '.xlsx':
                import openpyxl
                workbook = openpyxl.load_workbook(str(doc_path))
                quality_report['metrics']['worksheet_count'] = len(workbook.worksheets)
                workbook.close()
                
            elif doc_path.suffix.lower() == '.pptx':
                from pptx import Presentation
                presentation = Presentation(str(doc_path))
                quality_report['metrics']['slide_count'] = len(presentation.slides)
                
        except Exception as e:
            quality_report['issues'].append(f"Office document validation failed: {str(e)}")
    
    def _validate_text_file(self, text_path: Path, quality_report: Dict[str, Any]):
        """Validate text file."""
        try:
            with open(text_path, 'r', encoding='utf-8') as f:
                content = f.read()
                quality_report['metrics']['character_count'] = len(content)
                quality_report['metrics']['line_count'] = len(content.splitlines())
                
                if len(content.strip()) == 0:
                    quality_report['issues'].append("Text file is empty")
                    
        except Exception as e:
            quality_report['issues'].append(f"Text file validation failed: {str(e)}")


class LibreOfficeConverter:
    """LibreOffice headless converter."""
    
    def __init__(self, libreoffice_path: Optional[str] = None, timeout: int = 300):
        self.libreoffice_path = libreoffice_path or self._find_libreoffice()
        self.timeout = timeout
        self.quality_checker = ConversionQualityChecker()
        
        # Format mappings for LibreOffice
        self.format_filters = {
            ConversionFormat.PDF: {
                'docx': 'writer_pdf_Export',
                'xlsx': 'calc_pdf_Export', 
                'pptx': 'impress_pdf_Export',
                'odt': 'writer_pdf_Export',
                'ods': 'calc_pdf_Export',
                'odp': 'impress_pdf_Export',
                'rtf': 'writer_pdf_Export',
                'txt': 'writer_pdf_Export'
            },
            ConversionFormat.DOCX: {
                'pdf': 'MS Word 2007 XML',
                'odt': 'MS Word 2007 XML',
                'rtf': 'MS Word 2007 XML',
                'txt': 'MS Word 2007 XML'
            },
            ConversionFormat.ODT: {
                'docx': 'writer8',
                'pdf': 'writer8',
                'rtf': 'writer8',
                'txt': 'writer8'
            },
            ConversionFormat.TXT: {
                'docx': 'Text (encoded)',
                'pdf': 'Text (encoded)',
                'odt': 'Text (encoded)',
                'rtf': 'Text (encoded)'
            }
        }
    
    def _find_libreoffice(self) -> str:
        """Find LibreOffice executable."""
        possible_paths = [
            'libreoffice',
            '/usr/bin/libreoffice',
            '/opt/libreoffice/program/soffice',
            'C:\\Program Files\\LibreOffice\\program\\soffice.exe',
            'C:\\Program Files (x86)\\LibreOffice\\program\\soffice.exe'
        ]
        
        for path in possible_paths:
            if shutil.which(path):
                return path
                
        raise DocumentProcessingError("LibreOffice not found. Please install LibreOffice.")
    
    def convert_document(self, input_path: Path, output_path: Path, 
                        target_format: ConversionFormat, options: Dict[str, Any] = None) -> ConversionJob:
        """Convert a single document."""
        job = ConversionJob(
            input_path=input_path,
            output_path=output_path,
            target_format=target_format,
            options=options or {}
        )
        
        try:
            job.start_time = time.time()
            job.status = "converting"
            
            # Validate input file
            if not input_path.exists():
                raise DocumentProcessingError(f"Input file does not exist: {input_path}")
            
            # Create output directory if needed
            output_path.parent.mkdir(parents=True, exist_ok=True)
            
            # Get conversion filter
            input_format = input_path.suffix[1:].lower()  # Remove dot
            filter_name = self._get_conversion_filter(input_format, target_format)
            
            # Prepare LibreOffice command
            cmd = self._build_conversion_command(input_path, output_path, filter_name, options)
            
            # Execute conversion
            logger.info(f"Converting {input_path} to {target_format.value}")
            result = subprocess.run(
                cmd,
                capture_output=True,
                text=True,
                timeout=self.timeout,
                cwd=output_path.parent
            )
            
            if result.returncode != 0:
                raise DocumentProcessingError(f"LibreOffice conversion failed: {result.stderr}")
            
            # Check if output file was created
            if not output_path.exists():
                # LibreOffice might have created file with different name
                output_path = self._find_converted_file(output_path, target_format)
                if not output_path.exists():
                    raise DocumentProcessingError("Conversion completed but output file not found")
            
            job.end_time = time.time()
            job.status = "completed"
            
            # Perform quality check
            quality_report = self.quality_checker.check_conversion_quality(
                input_path, output_path, target_format
            )
            
            if not quality_report['is_valid']:
                job.status = "completed_with_issues"
                job.error_message = f"Quality issues: {', '.join(quality_report['issues'])}"
            
            logger.info(f"Conversion completed: {input_path} -> {output_path} ({job.duration:.2f}s)")
            
        except subprocess.TimeoutExpired:
            job.status = "failed"
            job.error_message = f"Conversion timeout after {self.timeout} seconds"
            job.end_time = time.time()
            
        except Exception as e:
            job.status = "failed"
            job.error_message = str(e)
            job.end_time = time.time()
            logger.error(f"Conversion failed: {e}")
        
        return job
    
    def _get_conversion_filter(self, input_format: str, target_format: ConversionFormat) -> str:
        """Get LibreOffice filter for conversion."""
        format_filters = self.format_filters.get(target_format, {})
        filter_name = format_filters.get(input_format)
        
        if not filter_name:
            # Try generic filters
            if target_format == ConversionFormat.PDF:
                if input_format in ['doc', 'docx', 'odt', 'rtf', 'txt']:
                    return 'writer_pdf_Export'
                elif input_format in ['xls', 'xlsx', 'ods']:
                    return 'calc_pdf_Export'
                elif input_format in ['ppt', 'pptx', 'odp']:
                    return 'impress_pdf_Export'
            
            raise DocumentProcessingError(
                f"No conversion filter found for {input_format} -> {target_format.value}"
            )
        
        return filter_name
    
    def _build_conversion_command(self, input_path: Path, output_path: Path, 
                                filter_name: str, options: Dict[str, Any]) -> List[str]:
        """Build LibreOffice conversion command."""
        cmd = [
            self.libreoffice_path,
            '--headless',
            '--convert-to', f"{output_path.suffix[1:]}:{filter_name}",
            '--outdir', str(output_path.parent),
            str(input_path)
        ]
        
        # Add additional options
        if options.get('no_logo'):
            cmd.append('--nologo')
        
        if options.get('invisible'):
            cmd.append('--invisible')
            
        return cmd
    
    def _find_converted_file(self, expected_path: Path, target_format: ConversionFormat) -> Path:
        """Find the actual converted file (LibreOffice might change the name)."""
        # Check in the same directory with different extensions
        base_name = expected_path.stem
        directory = expected_path.parent
        
        possible_extensions = [f".{target_format.value}"]
        if target_format == ConversionFormat.PDF:
            possible_extensions.append('.pdf')
        
        for ext in possible_extensions:
            candidate = directory / f"{base_name}{ext}"
            if candidate.exists():
                return candidate
        
        return expected_path


class DocumentConversionService:
    """High-level document conversion service with batch processing."""
    
    def __init__(self, max_workers: int = 4, libreoffice_path: Optional[str] = None):
        self.converter = LibreOfficeConverter(libreoffice_path)
        self.max_workers = max_workers
        
    def convert_single_document(self, input_path: Path, output_path: Path,
                              target_format: ConversionFormat, options: Dict[str, Any] = None) -> ConversionJob:
        """Convert a single document."""
        return self.converter.convert_document(input_path, output_path, target_format, options)
    
    def convert_batch(self, conversion_jobs: List[Tuple[Path, Path, ConversionFormat, Dict[str, Any]]]) -> List[ConversionJob]:
        """Convert multiple documents in parallel."""
        jobs = []
        
        with ThreadPoolExecutor(max_workers=self.max_workers) as executor:
            # Submit all jobs
            future_to_job = {}
            for input_path, output_path, target_format, options in conversion_jobs:
                job = ConversionJob(
                    input_path=input_path,
                    output_path=output_path,
                    target_format=target_format,
                    options=options or {}
                )
                future = executor.submit(
                    self.converter.convert_document,
                    input_path, output_path, target_format, options
                )
                future_to_job[future] = job
                jobs.append(job)
            
            # Collect results
            completed_jobs = []
            for future in as_completed(future_to_job):
                completed_job = future.result()
                completed_jobs.append(completed_job)
                
                # Update progress
                logger.info(f"Batch progress: {len(completed_jobs)}/{len(jobs)} completed")
        
        return completed_jobs
    
    def convert_directory(self, input_dir: Path, output_dir: Path, 
                         target_format: ConversionFormat, 
                         file_patterns: List[str] = None,
                         options: Dict[str, Any] = None) -> List[ConversionJob]:
        """Convert all matching files in a directory."""
        if file_patterns is None:
            file_patterns = ['*.docx', '*.xlsx', '*.pptx', '*.odt', '*.ods', '*.odp', '*.rtf']
        
        # Find all matching files
        input_files = []
        for pattern in file_patterns:
            input_files.extend(input_dir.glob(pattern))
        
        # Prepare conversion jobs
        conversion_jobs = []
        for input_file in input_files:
            # Generate output path
            output_file = output_dir / f"{input_file.stem}.{target_format.value}"
            conversion_jobs.append((input_file, output_file, target_format, options))
        
        logger.info(f"Converting {len(conversion_jobs)} files from {input_dir} to {output_dir}")
        
        return self.convert_batch(conversion_jobs)
    
    def get_supported_formats(self) -> Dict[str, List[str]]:
        """Get supported input and output formats."""
        return {
            'input_formats': [
                'docx', 'xlsx', 'pptx',  # Microsoft Office
                'odt', 'ods', 'odp',     # LibreOffice
                'rtf', 'txt',            # Text formats
                'doc', 'xls', 'ppt'      # Legacy Office (if LibreOffice supports)
            ],
            'output_formats': [format.value for format in ConversionFormat]
        }
    
    def estimate_conversion_time(self, file_size_mb: float, target_format: ConversionFormat) -> float:
        """Estimate conversion time in seconds based on file size."""
        # Base time estimates (seconds per MB)
        base_times = {
            ConversionFormat.PDF: 2.0,
            ConversionFormat.DOCX: 1.5,
            ConversionFormat.TXT: 0.5,
            ConversionFormat.HTML: 1.0
        }
        
        base_time = base_times.get(target_format, 2.0)
        return max(5.0, file_size_mb * base_time)  # Minimum 5 seconds
    
    def cleanup_temp_files(self, temp_dir: Path):
        """Clean up temporary files created during conversion."""
        try:
            if temp_dir.exists() and temp_dir.is_dir():
                shutil.rmtree(temp_dir)
                logger.info(f"Cleaned up temporary directory: {temp_dir}")
        except Exception as e:
            logger.warning(f"Failed to cleanup temp directory {temp_dir}: {e}")


# Convenience functions
def convert_to_pdf(input_path: Path, output_path: Path = None, options: Dict[str, Any] = None) -> ConversionJob:
    """Convert document to PDF."""
    if output_path is None:
        output_path = input_path.with_suffix('.pdf')
    
    service = DocumentConversionService()
    return service.convert_single_document(input_path, output_path, ConversionFormat.PDF, options)


def batch_convert_to_pdf(input_files: List[Path], output_dir: Path, options: Dict[str, Any] = None) -> List[ConversionJob]:
    """Batch convert documents to PDF."""
    service = DocumentConversionService()
    
    conversion_jobs = []
    for input_file in input_files:
        output_file = output_dir / f"{input_file.stem}.pdf"
        conversion_jobs.append((input_file, output_file, ConversionFormat.PDF, options))
    
    return service.convert_batch(conversion_jobs)